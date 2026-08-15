"""Demo deck generator — python-pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ACCENT = RGBColor(0x2E, 0x5E, 0xAA)   # biru
DARK = RGBColor(0x1F, 0x29, 0x37)     # abu gelap
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)    # background terang
GRAY = RGBColor(0x64, 0x6F, 0x7E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    """lines: list of (text, size, color, bold)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = "Calibri"
    return tb


# ── Slide 1: Title ──────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_bg(s, DARK)
add_rect(s, 0, Inches(5.1), SW, Inches(0.12), ACCENT)
add_text(s, Inches(1), Inches(2.2), Inches(11), Inches(1.5),
         [("Demo Slide Presentasi", 54, RGBColor(0xFF, 0xFF, 0xFF), True)])
add_text(s, Inches(1), Inches(3.6), Inches(11), Inches(0.8),
         [("Dibuat dengan python-pptx 🧪", 24, RGBColor(0xBF, 0xD3, 0xF0), False)])
add_text(s, Inches(1), Inches(5.6), Inches(11), Inches(0.6),
         [("Jimmy Neutron • 13 Agustus 2026", 16, GRAY, False)])

# ── Slide 2: Agenda ─────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_bg(s, LIGHT)
add_rect(s, 0, 0, Inches(0.18), SH, ACCENT)
add_text(s, Inches(0.8), Inches(0.6), Inches(10), Inches(1),
         [("Agenda", 36, DARK, True)])
agenda = [
    ("01", "Apa itu python-pptx"),
    ("02", "Kenapa pilih PPTX"),
    ("03", "Cara pakai & kustomisasi"),
]
y = Inches(2.0)
for num, t in agenda:
    add_text(s, Inches(1.2), y, Inches(0.9), Inches(0.6),
             [(num, 28, ACCENT, True)])
    add_text(s, Inches(2.2), y, Inches(9), Inches(0.6),
             [(t, 22, DARK, False)])
    y += Inches(1.0)

# ── Slide 3: Konten 1 ───────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_bg(s, LIGHT)
add_rect(s, 0, 0, Inches(0.18), SH, ACCENT)
add_text(s, Inches(0.8), Inches(0.6), Inches(10), Inches(1),
         [("Apa itu python-pptx?", 36, DARK, True)])
bullets = [
    "Library Python untuk membuat & memodifikasi file PowerPoint (.pptx)",
    "Full programatik: slide, teks, shape, warna, layout — semua dari kode",
    "Output standar OOXML, kompatibel dengan MS Office & Google Slides",
    "Cocok untuk generate deck otomatis dari data / template",
]
tb = add_text(s, Inches(1.2), Inches(2.0), Inches(10.5), Inches(4),
              [(b, 20, DARK, False) for b in bullets])
for p in tb.text_frame.paragraphs:
    p.space_after = Pt(14)

# ── Slide 4: Konten 2 ───────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_bg(s, LIGHT)
add_rect(s, 0, 0, Inches(0.18), SH, ACCENT)
add_text(s, Inches(0.8), Inches(0.6), Inches(10), Inches(1),
         [("Kenapa pilih PPTX?", 36, DARK, True)])
items = [
    ("Bisa diedit ulang", "— buka di PowerPoint/Google Slides, tinggal ubah"),
    ("Format universal", "— dikirim ke siapa pun, dibuka di mana pun"),
    ("Konsisten & rapi", "— desain seragam, gak ada slide miring/berantakan"),
    ("Skalabel", "— 5 slide atau 50 slide, effort-nya sama"),
]
y = Inches(2.0)
for head, desc in items:
    add_text(s, Inches(1.2), y, Inches(10), Inches(0.6),
             [(head + " " + desc, 20, DARK, False)])
    y += Inches(1.0)

# ── Slide 5: Penutup ────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_bg(s, DARK)
add_rect(s, 0, Inches(5.1), SW, Inches(0.12), ACCENT)
add_text(s, Inches(1), Inches(2.6), Inches(11), Inches(1.2),
         [("Mau bikin deck beneran?", 44, RGBColor(0xFF, 0xFF, 0xFF), True)],
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.0), Inches(11), Inches(0.8),
         [("Tinggal kasih topik + jumlah slide → langsung jadi 🚀", 22, RGBColor(0xBF, 0xD3, 0xF0), False)],
         align=PP_ALIGN.CENTER)

OUT = "/home/jimmy-bot/.openclaw/workspace/slides/demo_deck.pptx"
prs.save(OUT)
print("saved:", OUT)
